"""Genera entrega/Informe-Taller-RabbitMQ.docx (máx. 4 páginas)."""
import pathlib

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RAIZ = pathlib.Path(__file__).resolve().parent.parent.parent
ENT = RAIZ / "entrega"
EV = RAIZ / "evidencias"
DG = ENT / "diagramas"
NARANJA = RGBColor(0xC2, 0x4E, 0x0E)
OSCURO = RGBColor(0x1F, 0x2A, 0x37)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)  # carta
for m in ("left_margin", "right_margin"):
    setattr(sec, m, Cm(1.8))
sec.top_margin = sec.bottom_margin = Cm(1.5)

st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(10)
st.paragraph_format.space_after = Pt(3)
st.paragraph_format.line_spacing = 1.0
for nivel, tam in ((1, 13), (2, 11)):
    h = doc.styles[f"Heading {nivel}"]
    h.font.name = "Calibri"
    h.font.size = Pt(tam)
    h.font.color.rgb = NARANJA if nivel == 1 else OSCURO
    h.paragraph_format.space_before = Pt(8 if nivel == 1 else 5)
    h.paragraph_format.space_after = Pt(3)


def sombrear(celda, color):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def parrafo(texto, negrita_ini=None, size=None, italic=False, align=None):
    p = doc.add_paragraph()
    if negrita_ini:
        r = p.add_run(negrita_ini)
        r.bold = True
        if size:
            r.font.size = Pt(size)
    r = p.add_run(texto)
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    if align:
        p.alignment = align
    return p


def tabla(filas, anchos, size=8.5, cabecera="1F2A37"):
    t = doc.add_table(rows=len(filas), cols=len(filas[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, fila in enumerate(filas):
        for j, val in enumerate(fila):
            c = t.cell(i, j)
            c.width = Cm(anchos[j])
            c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val)
            r.font.size = Pt(size)
            if i == 0:
                r.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                sombrear(c, cabecera)
            elif j == 0:
                r.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def imagen(ruta, ancho_cm, pie):
    doc.add_picture(str(ruta), width=Cm(ancho_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[-1].paragraph_format.space_after = Pt(0)
    parrafo(pie, size=8, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)


def imagenes_lado(rutas_pies, ancho_cm):
    t = doc.add_table(rows=2, cols=len(rutas_pies))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, (ruta, pie) in enumerate(rutas_pies):
        p = t.cell(0, j).paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(ruta), width=Cm(ancho_cm))
        q = t.cell(1, j).paragraphs[0]
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = q.add_run(pie)
        r.italic = True
        r.font.size = Pt(8)


# ---------------- Portada compacta ----------------
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("Arquitectura Event-Driven multilenguaje con RabbitMQ y Docker")
r.bold = True
r.font.size = Pt(16)
r.font.color.rgb = OSCURO
p = parrafo("Principios de Arquitectura de Software — Taller en grupo", align=WD_ALIGN_PARAGRAPH.CENTER, size=10)
p = parrafo("David Esteban Suarez Lozano (2221909) · Thomas Alejandro Perez Rojas (2221906)",
            "Integrantes: ", align=WD_ALIGN_PARAGRAPH.CENTER, size=9.5)
p.paragraph_format.space_after = Pt(0)
p = parrafo("Jesus David Santodomingo Carrascal (2221922)  —  Fecha: 28 de octubre de 2026",
            align=WD_ALIGN_PARAGRAPH.CENTER, size=9.5)
p.paragraph_format.space_after = Pt(0)
p = parrafo("https://github.com/Tapr16/taller-eventdriven-rabbitmq", "Repositorio: ", align=WD_ALIGN_PARAGRAPH.CENTER, size=9.5)

# ---------------- 1. Descripción ----------------
doc.add_heading("1. Descripción del sistema", 1)
parrafo("Simulamos los eventos de una tienda en línea. Cuatro productores generan eventos de dominio y cuatro "
        "consumidores reaccionan a ellos sin conocerse entre sí: el único punto de contacto es el broker RabbitMQ "
        "(AMQP 0-9-1). Cada servicio corre en su propio contenedor Docker y se compila dentro de su Dockerfile "
        "(multi-stage para Java y Go), por lo que solo se requiere Docker en cada nodo.")
tabla([
    ["Servicio", "Lenguaje / librería", "Rol", "Routing keys publicadas / bindings"],
    ["productor-ordenes", "Python / pika", "Productor", "orden.creada"],
    ["productor-pagos", "Java / amqp-client", "Productor", "pago.aprobado, pago.rechazado"],
    ["productor-inventario", "Node.js / amqplib", "Productor", "inventario.actualizado, inventario.agotado"],
    ["productor-envios", "Go / amqp091-go", "Productor", "envio.despachado, envio.entregado"],
    ["consumidor-notificaciones", "Python / pika", "Consumidor", "q.notificaciones ← orden.creada, envio.*"],
    ["consumidor-facturacion", "Java / amqp-client", "Consumidor", "q.facturacion ← pago.aprobado, pago.rechazado"],
    ["consumidor-auditoria (+UI web)", "Node.js / amqplib", "Consumidor", "q.auditoria ← # (todos los eventos)"],
    ["consumidor-analitica", "Go / amqp091-go", "Consumidor", "q.analitica ← orden.*, inventario.*, pago.*"],
], [4.4, 3.3, 2.0, 8.2])
parrafo("Un exchange de tipo topic y durable (tienda.eventos); la routing key es el tipo de evento (dominio.acción). "
        "Cada consumidor declara su propia cola durable y sus bindings con comodines (* = una palabra, # = cero o más), "
        "de modo que un mismo evento llega a varias colas (p. ej. orden.creada llega a notificaciones, auditoría y "
        "analítica): publish/subscribe con filtrado por tema. Los mensajes son JSON persistentes (delivery_mode=2) con un "
        "campo nodo que identifica la máquina de origen; los consumidores usan ack manual y prefetch. Todos los servicios "
        "reintentan la conexión si el broker aún no está disponible.", "Topología AMQP. ")

# ---------------- 2. Diagramas ----------------
doc.add_heading("2. Diagramas arquitecturales", 1)
imagen(DG / "escenario1.png", 15.5, "Figura 1. Escenario 1 — centralizado: broker, productores y consumidores en un solo nodo "
       "(1 docker-compose.yml; los servicios encuentran al broker por el nombre «rabbitmq» en la red interna de Docker).")
imagen(DG / "escenario2.png", 13.5, "Figura 2. Escenario 2 — 2 nodos: Nodo 1 con los 8 servicios y Nodo 2 con el broker "
       "(2 docker-compose.yml; la IP del broker se pasa por .env → BROKER_IP).")
imagen(DG / "escenario3.png", 15.5, "Figura 3. Escenario 3 — 3 nodos: Python+Java en el Nodo 1, Node.js+Go en el Nodo 2 y "
       "el broker en el Nodo 3 (3 docker-compose.yml). Los eventos entre nodos 1 y 2 siempre pasan por el broker.")

# ---------------- 3. Evidencias ----------------
doc.add_heading("3. Evidencias de ejecución", 1)
parrafo("Todas las capturas y logs completos están en la carpeta evidencias/ del repositorio (github.com/Tapr16/taller-eventdriven-rabbitmq). Se usaron dos PCs de la sala: "
        "PC A 10.6.62.18 y PC B 10.6.62.137.")
doc.add_heading("Escenario 1 — Centralizado", 2)
parrafo("docker compose ps muestra 9 contenedores (broker + 8 servicios). En el panel se ven 8 conexiones AMQP desde la red "
        "interna 172.18.0.x, el exchange tienda.eventos con sus 8 bindings y 4 colas durables con 1 consumidor y 0 mensajes "
        "pendientes. La tasa de entrega (~3 msg/s) supera a la de publicación (~1 msg/s) porque cada evento se copia a varias colas.")
imagenes_lado([(EV / "escenario1" / "07-rabbitmq-exchange-bindings.jpg", "Exchange topic y bindings"),
               (EV / "escenario1" / "08-rabbitmq-queues.jpg", "4 colas, 1 consumidor c/u"),
               (EV / "escenario1" / "09-ui-web-eventos.jpg", "UI web de auditoría (:8080)")], 5.4)
doc.add_heading("Escenario 2 — 2 nodos (PC B = Nodo 1 servicios, PC A = Nodo 2 broker)", 2)
parrafo("En el PC A solo corre esc2-rabbitmq y aparecen 8 conexiones TCP establecidas al puerto 5672 desde 10.6.62.137. "
        "El panel identifica cada conexión como servicio@esc2-nodo1, y en el PC B los 4 consumidores procesan los eventos:")
imagen(EV / "escenario2" / "04-logs-consumidores.jpg", 17.5, "Logs de los 4 consumidores en el Nodo 1 (PC B) recibiendo eventos a través del broker remoto.")
imagenes_lado([(EV / "escenario2" / "pcA-04-terminal-conexiones-desde-pcB.jpg", "PC A: conexiones TCP desde 10.6.62.137"),
               (EV / "escenario2" / "pcA-05-rabbitmq-connections.jpg", "Panel: 8 conexiones @esc2-nodo1")], 6.5)
doc.add_heading("Escenario 3 — 3 nodos (PC A = Nodo 3 broker + Nodo 1, PC B = Nodo 2)", 2)
parrafo("Por disponibilidad de equipos, el Nodo 1 y el broker compartieron máquina física, pero como proyectos compose "
        "separados y sin red Docker en común: el Nodo 1 se conecta por la IP de red (BROKER_IP=10.6.62.18), igual que el Nodo 2. "
        "El panel muestra 4 conexiones @esc3-nodo1 y 4 @esc3-nodo2, y los logs evidencian tráfico cruzado entre nodos:")
p = doc.add_paragraph()
r = p.add_run("[consumidor-go@esc3-nodo2] <- pago.aprobado de productor-java@esc3-nodo1 | ordenes=1 pagos_ok=1 ventas=977.22\n"
              "[consumidor-python@esc3-nodo1] <- envio.entregado de productor-go@esc3-nodo2 | SMS: la orden #9218 fue entregada")
r.font.name = "Consolas"
r.font.size = Pt(7.5)
imagenes_lado([(EV / "escenario3" / "pcA-04-rabbitmq-connections.jpg", "Panel: conexiones @esc3-nodo1 y @esc3-nodo2"),
               (EV / "escenario3" / "02-ui-web-eventos.png", "UI web en Nodo 2: eventos de ambos nodos")], 8.0)

# ---------------- 4. Comparación ----------------
doc.add_heading("4. Comparación entre escenarios", 1).paragraph_format.page_break_before = True
tabla([
    ["Criterio", "Esc. 1 — Centralizado", "Esc. 2 — 2 nodos", "Esc. 3 — 3 nodos"],
    ["Configuración", "Mínima: un compose; resolución por nombre de servicio (rabbitmq)",
     "Hay que conocer la IP del broker (.env) y abrir el puerto 5672", "Igual que el 2, más coordinación entre equipos"],
    ["Latencia", "Muy baja (red virtual local)", "Añade un salto de red LAN", "Igual que el 2; los eventos Nodo 1↔Nodo 2 pasan siempre por el Nodo 3"],
    ["Escalabilidad", "Limitada a los recursos de una máquina", "El broker no compite por CPU/RAM con los servicios",
     "Servicios repartidos; se agregan nodos sin tocar los existentes"],
    ["Tolerancia a fallos", "Cae el nodo → cae todo",
     "Cae el Nodo 1 → los mensajes esperan en colas durables; cae el Nodo 2 → nadie se comunica",
     "Cae el Nodo 1 o 2 → el otro sigue operando; el broker sigue siendo punto único de falla"],
    ["Seguridad", "Puertos solo locales", "AMQP expuesto en la LAN: usuario propio (no guest); idealmente TLS", "Igual que el 2"],
    ["Depuración", "Todos los logs en un solo lugar", "Logs repartidos en 2 máquinas", "Logs en 3 máquinas; el campo nodo y el nombre de conexión ayudan a rastrear"],
], [2.6, 4.4, 5.3, 5.6], size=8)

parrafo("desacoplamiento total (movimos consumidores entre nodos sin modificar productores ni cambiar código, solo los "
        "docker-compose), interoperabilidad entre 4 lenguajes gracias a AMQP + JSON, y amortiguación: si un consumidor se "
        "detiene, sus mensajes se acumulan en la cola y se procesan al volver.", "Ventajas del estilo event-driven: ")
parrafo("el broker es un punto único de falla (se mitigaría con un clúster RabbitMQ y quorum queues), la consistencia es "
        "eventual y es más difícil seguir un flujo de punta a punta sin trazas distribuidas.", "Limitaciones: ")
doc.add_heading("Retos técnicos encontrados", 2)
for b in [
    "El usuario guest de RabbitMQ solo acepta conexiones locales → se creó el usuario admin para los escenarios distribuidos.",
    "Los servicios pueden arrancar antes que el broker → reintento de conexión en los 4 lenguajes, y healthcheck + depends_on en el escenario 1.",
    "Docker Desktop publica el puerto con NAT, así que RabbitMQ ve a todos los clientes remotos como 172.18.0.1 → se nombró cada "
    "conexión AMQP (servicio@nodo) y se verificaron las IP reales con Get-NetTCPConnection en el host del broker.",
    "Firewall de Windows y red de la universidad al exponer el 5672 entre PCs; Docker Desktop fallaba en los PCs compartidos por "
    "sockets de sesiones anteriores (script arreglar-docker.ps1).",
    "Compilar Java y Go sin tenerlos instalados → multi-stage builds en los Dockerfile.",
]:
    p = doc.add_paragraph(b, style="List Bullet")
    p.paragraph_format.space_after = Pt(1)

# ---------------- 5. Conclusiones ----------------
doc.add_heading("5. Conclusiones", 1)
for b in [
    "El mismo código y las mismas imágenes funcionaron sin cambios en los tres escenarios: solo cambió dónde corre cada contenedor "
    "y la dirección del broker. Esa es la principal ventaja del estilo event-driven con un broker: la topología física queda fuera del código.",
    "El exchange topic con bindings por comodines permitió que un mismo evento llegue a varios consumidores con intereses distintos, "
    "sin que el productor sepa quién lo consume.",
    "Distribuir mejora el aislamiento de recursos y la tolerancia a la caída de nodos de servicios, a cambio de más configuración de red, "
    "seguridad y dificultad de depuración; el broker centralizado sigue siendo el punto crítico, por lo que en producción se usaría un clúster.",
]:
    p = doc.add_paragraph(b, style="List Bullet")
    p.paragraph_format.space_after = Pt(1)

out = ENT / "Informe-Taller-RabbitMQ.docx"
doc.save(out)
print(out)
